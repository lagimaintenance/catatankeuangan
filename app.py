"""ConvertRK multi-bank: PDF rekening koran menjadi satu sheet GL berformat asli.
Jalankan: python -m streamlit run app.py
"""
from __future__ import annotations
from datetime import date
from copy import copy
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import sys
import platform
import pdfplumber
import xml.etree.ElementTree as ET
import zipfile
import hashlib

import streamlit as st
from openpyxl import load_workbook
from multi_bank_parser import ParsingError, parse_bank_pdf, expand_uploads, identify

GagalBaca = ParsingError

# ====== 2. MENGISI TEMPLATE EXCEL ASLI (FORMAT BANK) ======
# Bank_template.xlsx adalah versi KOSONG dan anonim dari file Bank(1).xlsx.
# Font, warna, border, tinggi header, dan lebar kolom tetap dari file contoh.
# Semua transaksi masuk ke SATU sheet GL, hanya 10 kolom A:J.

TEMPLATE = Path(__file__).with_name('Bank_template.xlsx')
KOLOM = [
    'Tanggal', 'No. Bukti', 'Kode Akun', 'Kode Lawan Akun', 'Nama Akun',
    'Uraian', 'Debet', 'Kredit', 'Saldo', 'Nama Vendor',
]


def teks_aman_excel(value) -> str:
    """Singkirkan karakter XML ilegal, agar hasil XLSX dapat dibuka Excel."""
    text = str(value)
    text = ''.join(
        ch for ch in text
        if (ord(ch) in (9, 10, 13) or 32 <= ord(ch) <= 0xD7FF
            or 0xE000 <= ord(ch) <= 0xFFFD or 0x10000 <= ord(ch) <= 0x10FFFF)
        and (ord(ch) & 0xFFFF) not in (0xFFFE, 0xFFFF)
    )[:32767]
    # Penting: jangan izinkan keterangan dari PDF menjadi rumus Excel.
    if text.startswith(('=', '+', '-', '@')):
        text = "'" + text[:32766]
    return text


def buat_excel(statements, company: str, code: str) -> bytes:
    """Konversi PDF ke 1 sheet GL menggunakan style ASLI Bank(1).xlsx."""
    if not TEMPLATE.is_file():
        raise GagalBaca(
            'Bank_template.xlsx tidak ditemukan. Pastikan file ini berada '
            'di folder yang sama dengan app.py (termasuk saat deploy).'
        )
    try:
        book = load_workbook(TEMPLATE, keep_links=False)
        if len(book.worksheets) != 1:
            raise GagalBaca('Template harus berisi hanya satu sheet.')
        ws = book.active
        ws.title = 'GL'
        header = [ws.cell(5, col).value for col in range(1, 11)]
        if header != KOLOM:
            raise GagalBaca('Header Bank_template.xlsx tidak sesuai contoh Buku Bank.')

        # Salin style asli sebelum membuat ulang data. _style berisi seluruh
        # atribut font, color, border, alignment, protection, dan num format.
        saldo_style = [copy(ws.cell(7, col)._style) for col in range(1, 11)]
        mutasi_style = [copy(ws.cell(8, col)._style) for col in range(1, 11)]
        original_row7_height = ws.row_dimensions[7].height
        original_row8_height = ws.row_dimensions[8].height

        # Isi lama template memang kosong, namun bersihkan kembali agar
        # jika template diganti pengguna tidak ada transaksi yang tersisa.
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row, max_col=10):
            for cell in row:
                cell.value = None
        ws['A1'] = teks_aman_excel(company.strip() or 'BUKU BANK')
        ws['A2'] = 'Buku Bank'
        year_set = sorted({s['year'] for s in statements})
        ws['A3'] = year_set[0] if len(year_set) == 1 else ', '.join(map(str, year_set))

        bank_code = (int(code) if code.isdigit() and len(code) <= 14
                     and not code.startswith('0') else teks_aman_excel(code))
        row_num = 7
        statements = sorted(statements, key=lambda s: (
            s['rekening'], s['year'], s['month'], s['file']))
        for index, stmt in enumerate(statements):
            # Satu baris SALDO BANK AWAL untuk tiap periode/rekening.
            # Bukan mutasi, sehingga tidak dihitung sebagai debit/kredit.
            for col in range(1, 11):
                ws.cell(row_num, col)._style = copy(saldo_style[col - 1])
            ws.cell(row_num, 6).value = f'SALDO BANK AWAL - {stmt["bank"]} - {stmt["rekening"]} - {stmt["month"]:02d}/{stmt["year"]}'
            ws.cell(row_num, 9).value = float(stmt['saldo_awal'])
            ws.row_dimensions[row_num].height = original_row7_height
            row_num += 1

            for item in stmt['transactions']:
                values = [
                    item['tanggal'], None, bank_code, None, None,
                    teks_aman_excel(item['uraian']),
                    float(item['debet']) if item['debet'] else None,
                    float(item['kredit']) if item['kredit'] else None,
                    float(item['saldo']), None,
                ]
                for col, value in enumerate(values, 1):
                    cell = ws.cell(row_num, col)
                    cell._style = copy(mutasi_style[col - 1])
                    cell.value = value
                # Multi-baris keterangan tetap lengkap. Styling tidak berubah.
                lines = item['uraian'].count('\n') + 1
                ws.row_dimensions[row_num].height = (
                    min(110, max(17, 14 * lines + 4)) if lines > 1
                    else original_row8_height
                )
                row_num += 1

        # Hapus area sisa kosong dari template agar satu-satunya tabel
        # berakhir di transaksi terakhir; tidak ada data bank contoh.
        if ws.max_row >= row_num:
            ws.delete_rows(row_num, ws.max_row - row_num + 1)
        # Header, font, border, format tanggal/angka, dan lebar kolom
        # sengaja TIDAK diubah agar sesuai persis dengan contoh.

        buffer = BytesIO()
        book.save(buffer)
        result = buffer.getvalue()
        book.close()

        # Pemeriksaan internal struktur agar Excel tidak meminta repair.
        with zipfile.ZipFile(BytesIO(result)) as archive:
            broken = archive.testzip()
            if broken:
                raise GagalBaca(f'File XLSX korup: {broken}')
            for name in archive.namelist():
                if name.endswith(('.xml', '.rels')):
                    ET.fromstring(archive.read(name))
        checked = load_workbook(BytesIO(result), read_only=True, keep_links=False)
        try:
            if checked.sheetnames != ['GL']:
                raise GagalBaca('Output bukan satu sheet GL.')
            if [checked['GL'].cell(5, c).value for c in range(1, 11)] != KOLOM:
                raise GagalBaca('Header output berubah.')
            if checked['GL'].max_row != row_num - 1:
                raise GagalBaca('Jumlah baris output tidak sesuai transaksi.')
        finally:
            checked.close()
        return result
    except GagalBaca:
        raise
    except Exception as exc:
        raise GagalBaca(f'Gagal membuat Excel: {type(exc).__name__}: {exc}') from exc



def diagnostic_pdf(content: bytes, name: str) -> dict:
    """Metadata teknis saja, TANPA membocorkan isi dan nomor rekening."""
    details={'Nama file':name, 'Python':platform.python_version(),
             'pdfplumber':pdfplumber.__version__}
    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            first=pdf.pages[0].extract_text() or ''
            details.update({'Halaman':len(pdf.pages),
                            'Karakter teks halaman 1':len(first),
                            'Bank terbaca':identify(first),
                            'Ada MUTASI': 'MUTASI' in first.upper(),
                            'Ada SALDO': 'SALDO' in first.upper(),
                            'Ada REKENING GIRO': 'REKENING GIRO' in first.upper()})
    except Exception as exc:
        details['Kesalahan pembacaan']=f'{type(exc).__name__}: {exc}'
    return details


# ===== STREAMLIT =====
def rupiah(n):
    return 'Rp ' + f'{n:,.2f}'.replace(',', 'X').replace('.', ',').replace('X','.')


def main():
    st.set_page_config(page_title='ConvertRK Multi-Bank',page_icon='📑',layout='wide')
    st.title('ConvertRK — PDF Rekening Koran ke GL')
    st.caption('Bank BCA, BSI, CIMB Niaga, Mandiri, Muamalat, Permata, serta bank lain dengan tabel PDF standar. ' 
               'Satu sheet GL, format Buku Bank asli, tanpa mapping akun.')
    a,b=st.columns([3,1])
    with a:
        company=st.text_input('Nama perusahaan (opsional)',value='BUKU BANK')
    with b:
        code=st.text_input('Kode akun bank (untuk seluruh mutasi)',value='111',help='Kode Lawan Akun dan Nama Akun tetap kosong.')
    uploads=st.file_uploader('Upload satu atau beberapa PDF/ZIP rekening koran',
                             type=['pdf','zip'],accept_multiple_files=True)
    if not uploads:
        st.info('Gunakan PDF digital berbasis teks. PDF scan/gambar belum didukung. ' 
                'Bank lain akan dicoba menggunakan pembacaan tabel standar.')
        return
    signature=(tuple((f.name,f.size,hashlib.sha256(f.getvalue()).hexdigest()) for f in uploads),company,code)
    if st.button('Konversi ke GL',type='primary'):
        st.session_state.pop('convertRK_result',None)
        try:
            candidates=expand_uploads(uploads)
            if not candidates:raise ParsingError('Tidak ditemukan PDF dalam unggahan.')
            statements=[];failures=[];dupes=[];seen=set(); diagnostics=[]
            progress=st.progress(0,text='Memproses PDF rekening koran...')
            for index,(filename,content) in enumerate(candidates,1):
                sha=hashlib.sha256(content).hexdigest()
                if sha in seen:
                    dupes.append(filename + ': file identik sudah diproses, tidak dimasukkan dua kali.')
                    continue
                seen.add(sha)
                progress.progress(index/len(candidates),text=f'Membaca {filename}... ({index}/{len(candidates)})')
                try:statements.append(parse_bank_pdf(content,filename))
                except ParsingError as err:
                    failures.append(f'{filename}: {err}')
                    diagnostics.append(diagnostic_pdf(content,filename))
            progress.empty()
            if not statements:
                if diagnostics:
                    with st.expander('Diagnostik parser Cloud (tanpa isi rekening koran)', expanded=True):
                        st.dataframe(diagnostics,hide_index=True,use_container_width=True)
                raise ParsingError('Tidak ada PDF berhasil diproses.\n'+'\n'.join(failures))
            output=buat_excel(statements,company,code.strip() or '111')
            st.session_state['convertRK_result']=(signature,statements,output,failures,dupes,diagnostics)
        except ParsingError as err:st.error(str(err))
    result=st.session_state.get('convertRK_result')
    if not result:return
    saved,statements,data,failures,dupes,diagnostics=result
    if saved!=signature:
        st.warning('Unggahan/pengaturan berubah. Klik Konversi lagi.')
        return
    m1,m2,m3=st.columns(3)
    m1.metric('Transaksi dibaca',f'{sum(len(x["transactions"]) for x in statements):,}')
    m2.metric('Pemasukan / Debet GL',rupiah(sum((x['debet'] for x in statements),Decimal(0))))
    m3.metric('Pengeluaran / Kredit GL',rupiah(sum((x['kredit'] for x in statements),Decimal(0))))
    if failures:
        st.error(f'{len(failures)} PDF tidak berhasil diproses — data dari PDF ini TIDAK masuk Excel:')
        for f in failures:st.write('• '+f)
        if diagnostics:
            with st.expander('Diagnostik PDF gagal (tanpa isi rekening koran)', expanded=True):
                st.dataframe(diagnostics,hide_index=True,use_container_width=True)
    if dupes:
        for d in dupes:st.info(d)
    issues=[(s['file'],x) for s in statements for x in s['warnings']]
    critical=lambda x:not ('dihitung dari transaksi pertama' in x or 'memakai saldo transaksi terakhir' in x)
    if any(critical(msg) for _,msg in issues):
        st.error('Sebagian transaksi belum dapat dinyatakan lengkap/cocok. Periksa catatan rekonsiliasi sebelum memakai GL untuk pembukuan.')
    elif not failures:
        st.success('Seluruh file yang dibaca lolos pemeriksaan yang tersedia (catatan inferensi saldo ditampilkan jika ada).')
    if issues:
        with st.expander(f'Catatan dan peringatan: {len(issues)}',expanded=True):
            for file,msg in issues:st.write(f'• **{file}**: {msg}')
    st.subheader('Rekap PDF')
    st.dataframe([{'Bank':s['bank'],'File':s['file'],'Rekening':s['rekening'],
                   'Periode':f'{s["month"]:02d}/{s["year"]}',
                   'Transaksi':len(s['transactions']),
                   'Saldo Awal':float(s['saldo_awal']),'Masuk (Debet GL)':float(s['debet']),
                   'Keluar (Kredit GL)':float(s['kredit']), 'Saldo Akhir':float(s['saldo_akhir']),
                   'Status':'SESUAI' if not s['warnings'] else 'CATATAN/PERLU CEK'} for s in statements],
                 hide_index=True,use_container_width=True)
    sample=[]
    for s in statements:
        for t in s['transactions'][:10]:
            sample.append({'Bank':s['bank'],'Tanggal':str(t['tanggal']),'Uraian':t['uraian'],
                           'Debet':float(t['debet']),'Kredit':float(t['kredit']),'Saldo':float(t['saldo'])})
    with st.expander('Pratinjau 10 transaksi pertama per rekening',expanded=False):
        st.dataframe(sample,hide_index=True,use_container_width=True)
    st.download_button('⬇️ Download GL Excel (1 Sheet, Format Asli)',data=data,
                       file_name='GL_Rekening_Koran.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
    st.caption('Mapping akun dikosongkan untuk pemeriksaan manual. ' 
               'Untuk data bank rahasia, jalankan aplikasi di Mac lokal. ' 
               'Jangan mengunggah PDF rekening koran ke GitHub.')

if __name__ == '__main__':main()
